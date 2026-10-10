import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
COLOR_PRIMARY = RGBColor(63, 81, 181)      # Indigo
COLOR_SECONDARY = RGBColor(26, 35, 126)   # Dark Indigo
COLOR_TEXT = RGBColor(33, 33, 33)         # Dark Gray
COLOR_MUTED = RGBColor(108, 117, 125)     # Gray

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(12)
    run.font.color.rgb = COLOR_MUTED
    return p

def add_section_heading(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f'{number}. {text}')
    run.font.name = 'Segoe UI'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_subsection_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_body_p(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Segoe UI'
        r_prefix.font.size = Pt(10.5)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = COLOR_TEXT
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_TEXT
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Segoe UI'
        r_prefix.font.size = Pt(10.5)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = COLOR_TEXT
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_TEXT
    return p

# TITLE
add_header_title(doc, 'DOCUMENTO DE VISÃO DO NEGÓCIO')
add_subtitle(doc, 'Sistema E-Commerce Loja Games (Framework SPA Angular)\nFramework_Consumo_API_SP2 - Documentação de Sistema')

meta_tbl = doc.add_table(rows=4, cols=2)
meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ('Projeto:', 'Loja Games (E-commerce Web SPA)'),
    ('Versão da Aplicação:', '1.0.0 (Angular 15 / Material / Bootstrap)'),
    ('Responsável Técnico:', 'Equipe de Desenvolvimento de Software'),
    ('Data de Emissão:', 'Outubro de 2026')
]
for i, (k, v) in enumerate(meta_data):
    cell_k, cell_v = meta_tbl.rows[i].cells
    cell_k.text = k
    cell_v.text = v
    set_cell_background(cell_k, 'E8EAF6')
    set_cell_background(cell_v, 'F5F7FB')
    cell_k.paragraphs[0].runs[0].font.bold = True
    cell_k.paragraphs[0].runs[0].font.color.rgb = COLOR_PRIMARY
    cell_k.paragraphs[0].runs[0].font.size = Pt(9.5)
    cell_v.paragraphs[0].runs[0].font.size = Pt(9.5)
    set_cell_margins(cell_k, 50, 50, 100, 100)
    set_cell_margins(cell_v, 50, 50, 100, 100)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# 1. FINALIDADE
add_section_heading(doc, 1, 'Finalidade')
add_body_p(doc, 'Este documento tem como finalidade estabelecer a visão estratégica, arquitetural e funcional do sistema Loja Games, definindo com clareza seus objetivos de negócio, público-alvo, requisitos essenciais e fluxo de operações. Ele atua como guia normativo tanto para usuários finais — que buscam compreender os recursos e a navegabilidade do sistema — quanto para engenheiros e programadores, auxiliando nas rotinas de manutenção, testes e futuras implementações (como o consumo pleno de APIs REST e microsserviços de pagamento).')

# 2. ESCOPO
add_section_heading(doc, 2, 'Escopo')
add_body_p(doc, 'O sistema Loja Games consiste em uma plataforma de comércio eletrônico focada no varejo digital de jogos eletrônicos. O escopo compreende:')
add_bullet(doc, 'Módulo de Navegação Superior (Menu): Exibição do logotipo institucional, acesso direto à vitrine de produtos e link para autenticação.', '• ')
add_bullet(doc, 'Módulo de Vitrine e Promoções (Home/Início): Apresentação de campanhas promocionais através de carrossel de banners e grade de cards com jogos em destaque, preços e botões de compra.', '• ')
add_bullet(doc, 'Módulo de Autenticação (Login): Interface limpa e segura para autenticação de clientes, permitindo o preenchimento de usuário e senha com validações locais e redirecionamento de rotas.', '• ')
add_bullet(doc, 'Módulo Rodapé (Footer): Informações institucionais, direitos autorais e créditos do projeto.', '• ')
add_body_p(doc, 'Ficam fora do escopo desta versão preliminar: a integração com gateways bancários externos para cobrança em tempo real e o controle de inventário físico, itens previstos para os próximos ciclos de desenvolvimento.')

# 3. VISÃO GERAL
add_section_heading(doc, 3, 'Visão Geral')
add_body_p(doc, 'A aplicação é desenvolvida como uma Single Page Application (SPA) baseada no ecossistema Angular 15. A arquitetura desacoplada possibilita alta velocidade de carregamento, dispensando recarregamentos completos de página (full reload) durante a navegação. O design visual foi concebido para o público gamer, integrando componentes visuais do Angular Material (botões, cards, campos de formulário, grid) com o sistema de grid e carrossel do Bootstrap 5.')

# 4. VISÃO GERAL DO PRODUTO
add_section_heading(doc, 4, 'Visão Geral do Produto')
add_body_p(doc, 'A Loja Games se diferencia por aliar uma estética moderna no estilo gamer neon a uma estrutura limpa e ergonômica. Seus principais atributos de produto incluem:')
add_bullet(doc, 'Responsividade Total: Adaptabilidade automática a desktops, notebooks, tablets e smartphones.', '• ')
add_bullet(doc, 'Arquitetura Modular: Componentização em TypeScript que facilita a substituição ou enriquecimento de módulos sem efeitos colaterais.', '• ')
add_bullet(doc, 'Prontidão para Consumo de API: Estrutura preparada para receber serviços Angular (HttpClient) conectados a APIs REST (ex.: json-server ou backend corporativo).', '• ')

# 5. POSICIONAMENTO NO MERCADO
add_section_heading(doc, 5, 'Posicionamento no Mercado')
add_body_p(doc, 'O mercado global e nacional de jogos eletrônicos tem apresentado crescimento exponencial, com migração massiva para formatos 100% digitais. Nesse cenário, a Loja Games posiciona-se como uma vitrine ágil e atrativa, focada no público jovem e entusiasta de games que busca acesso rápido às principais promoções sem fricções excessivas no fluxo de visualização e checkout.')

# 6. DESCRIÇÃO DA EQUIPE DESENVOLVEDORA E DOS CLIENTES
add_section_heading(doc, 6, 'Descrição da Equipe Desenvolvedora e dos Clientes')
add_subsection_heading(doc, 'Equipe de Desenvolvimento')
add_bullet(doc, 'Desenvolvedores Front-End: Responsáveis pela implementação dos componentes Angular, templates HTML, estilos CSS e regras de negócio no TypeScript.', '• ')
add_bullet(doc, 'UI/UX Designers: Encarregados da prototipação, harmonia de cores, experiência visual gamer e ergonomia das telas.', '• ')
add_bullet(doc, 'Engenheiro de QA / Testes: Responsável por validar os fluxos funcionais, integridade de rotas e compatibilidade entre navegadores.', '• ')
add_bullet(doc, 'Product Owner (PO): Define a priorização do backlog de requisitos e valida os entregáveis perante o modelo de negócio.', '• ')

add_subsection_heading(doc, 'Perfis de Clientes e Usuários')
add_bullet(doc, 'Gamer Visitante (Não Autenticado): Acessa a loja, visualiza novidades, inspeciona preços e detalhes dos jogos.', '• ')
add_bullet(doc, 'Cliente Cadastrado: Usuário autenticado que efetua a aquisição dos títulos digitais e gerencia seus dados.', '• ')
add_bullet(doc, 'Administrador de Catálogo: Usuário corporativo que atualiza os produtos, gerencia preços promocionais e controla os banners de destaque.', '• ')

# 7. ALTERNATIVAS E CONCORRÊNCIAS
add_section_heading(doc, 7, 'Alternativas e Concorrências')
add_body_p(doc, 'As principais plataformas concorrentes no ecossistema de distribuição digital incluem:')
add_bullet(doc, 'Steam (Valve): Líder consolidada, com catálogo massivo, mas interface com alta densidade de informações.', '• ')
add_bullet(doc, 'Epic Games Store: Foco em títulos gratuitos e promoções semanais agressivas.', '• ')
add_bullet(doc, 'Nuuvem: Plataforma brasileira com promoções locais e suporte regionalizado.', '• ')
add_body_p(doc, 'Diferencial Competitivo: A Loja Games foca em uma interface minimalista, rápida, livre de sobrecarga de dados desnecessários, com acesso imediato às melhores promoções em um único clique.')

# 8. RESTRIÇÕES
add_section_heading(doc, 8, 'Restrições')
add_bullet(doc, 'Plataforma e Framework: A aplicação deve ser construída obrigatoriamente sobre Angular 15 e TypeScript.', '• ')
add_bullet(doc, 'Ambiente de Execução: Suporte estrito aos principais navegadores modernos (Google Chrome, Mozilla Firefox, Microsoft Edge, Safari) com suporte a ECMAScript 6+.', '• ')
add_bullet(doc, 'Conformidade Legal: Conformidade com as diretrizes da Lei Geral de Proteção de Dados (LGPD - Lei nº 13.709/2018), resguardando credenciais e dados pessoais dos usuários.', '• ')
add_bullet(doc, 'Restrição de Conectividade: Operação client-side dependente de conexão com internet para download inicial dos bundles e consumo de recursos remotos.', '• ')

# 9. REQUISITOS DO PRODUTO
add_section_heading(doc, 9, 'Requisitos do Produto')

add_subsection_heading(doc, 'Requisitos Funcionais (RF)')
rf_table = doc.add_table(rows=7, cols=3)
rf_table.alignment = WD_TABLE_ALIGNMENT.CENTER
rf_headers = ['Identificador', 'Nome do Requisito', 'Descrição Funcional']
for col_idx, text in enumerate(rf_headers):
    cell = rf_table.rows[0].cells[col_idx]
    cell.text = text
    set_cell_background(cell, '3F51B5')
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    cell.paragraphs[0].runs[0].font.size = Pt(9.5)

rf_data = [
    ('RF01', 'Visualizar Catálogo de Jogos', 'O sistema deve exibir a grade de jogos com imagem, nome, descrição breve, preço promocional e botão de compra.'),
    ('RF02', 'Navegação por Banners', 'O sistema deve disponibilizar um carrossel rotativo interativo com os principais destaques promocionais.'),
    ('RF03', 'Autenticação de Usuário (Login)', 'O sistema deve permitir que o usuário digite suas credenciais (e-mail/usuário e senha) para acessar a plataforma.'),
    ('RF04', 'Ação de Compra de Título', 'O sistema deve disponibilizar o acionamento de compra direta em cada card de produto.'),
    ('RF05', 'Roteamento SPA Dinâmico', 'O sistema deve permitir transição instantânea entre as rotas /inicio e /login sem recarregamento da página.'),
    ('RF06', 'Retorno à Vitrine', 'A tela de login deve fornecer um atalho de retorno rápido para a página inicial.')
]
for row_idx, data in enumerate(rf_data, start=1):
    for col_idx, val in enumerate(data):
        cell = rf_table.rows[row_idx].cells[col_idx]
        cell.text = val
        cell.paragraphs[0].runs[0].font.size = Pt(9.0)
        set_cell_background(cell, 'F5F7FB' if row_idx % 2 == 1 else 'FFFFFF')
        set_cell_margins(cell, 45, 45, 70, 70)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

add_subsection_heading(doc, 'Requisitos Não Funcionais (RNF)')
rnf_table = doc.add_table(rows=6, cols=3)
rnf_table.alignment = WD_TABLE_ALIGNMENT.CENTER
rnf_headers = ['Identificador', 'Categoria', 'Descrição do Critério']
for col_idx, text in enumerate(rnf_headers):
    cell = rnf_table.rows[0].cells[col_idx]
    cell.text = text
    set_cell_background(cell, '1A237E')
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    cell.paragraphs[0].runs[0].font.size = Pt(9.5)

rnf_data = [
    ('RNF01', 'Desempenho', 'O tempo de resposta entre a navegação de rotas deve ser inferior a 1 segundo.'),
    ('RNF02', 'Responsividade', 'A interface deve ser totalmente responsiva para larguras de tela de 320px a 2560px.'),
    ('RNF03', 'Usabilidade', 'A interface deve seguir as diretrizes do Material Design, com contrastes adequados e tipografia clara.'),
    ('RNF04', 'Manutenibilidade', 'O código deve ser modularizado em componentes independentes com TypeScript tipado.'),
    ('RNF05', 'Segurança', 'O campo de senha deve mascarar os caracteres digitados para proteger contra visualizações indevidas.')
]
for row_idx, data in enumerate(rnf_data, start=1):
    for col_idx, val in enumerate(data):
        cell = rnf_table.rows[row_idx].cells[col_idx]
        cell.text = val
        cell.paragraphs[0].runs[0].font.size = Pt(9.0)
        set_cell_background(cell, 'F5F7FB' if row_idx % 2 == 1 else 'FFFFFF')
        set_cell_margins(cell, 45, 45, 70, 70)

# 10. MODELAGEM CASO DE USO
add_section_heading(doc, 10, 'Modelagem Caso de Uso')
add_body_p(doc, 'A modelagem de casos de uso representa as interações entre os atores (usuários) e o sistema, detalhando as rotinas funcionais e a navegabilidade geral da solução.')

# INSERT DIAGRAM IMAGE
if os.path.exists('diagrama_casos_de_uso.png'):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture('diagrama_casos_de_uso.png', width=Inches(6.0))
    
    p_leg = doc.add_paragraph()
    p_leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_leg.paragraph_format.space_after = Pt(12)
    r_leg = p_leg.add_run('Figura 1: Diagrama Visual de Casos de Uso do Sistema Loja Games')
    r_leg.font.size = Pt(9.0)
    r_leg.font.italic = True
    r_leg.font.color.rgb = COLOR_MUTED

add_subsection_heading(doc, 'Componentes da Modelagem')
add_bullet(doc, 'Entidades externas que interagem com o sistema (Visitante e Cliente Autenticado).', 'Ator: ')
add_bullet(doc, 'Funcionalidade ou serviço provido pelo sistema em benefício de um ator (UC01 a UC05).', 'Use Case: ')
add_bullet(doc, 'Sequência cronológica de eventos desencadeada pela ação do usuário perante o sistema.', 'Cenário: ')
add_bullet(doc, 'Ligação de comunicação e associação entre o ator e o caso de uso executado.', 'Comunicação: ')

add_subsection_heading(doc, 'Especificação Detalhada dos Casos de Uso')

ucs = [
    {
        'id': 'UC01',
        'nome': 'Visualizar Vitrine de Jogos em Promoção',
        'ator': 'Usuário / Gamer (Visitante) e Cliente Autenticado',
        'comunicacao': 'Acesso direto à rota raiz /inicio.',
        'cenario_principal': '1. O usuário acessa a plataforma; 2. O sistema carrega a lista de produtos em promoção com suas capas, títulos, descrições e valores; 3. O usuário visualiza as informações de forma clara.',
        'cenario_alternativo': 'Se houver falha de rede ou recurso indisponível, o sistema exibe mensagens ou imagens de fallback mantendo a estabilidade da interface.'
    },
    {
        'id': 'UC02',
        'nome': 'Navegar no Carrossel de Banners',
        'ator': 'Usuário / Gamer (Visitante) e Cliente Autenticado',
        'comunicacao': 'Interação com os botões "Anterior" e "Próximo" ou rotação automática.',
        'cenario_principal': '1. O usuário visualiza o banner promocional; 2. O usuário clica na seta de navegação; 3. O carrossel transiciona de forma suave para o próximo destaque promocional.',
        'cenario_alternativo': 'O carrossel realiza a transição temporizada automática caso não haja clique do usuário.'
    },
    {
        'id': 'UC03',
        'nome': 'Realizar Login / Autenticação',
        'ator': 'Usuário / Gamer (Visitante) e Cliente Autenticado',
        'comunicacao': 'Navegação para a rota /login e envio dos dados pelo botão "Entrar".',
        'cenario_principal': '1. O usuário clica em "Login" no menu superior; 2. O sistema exibe o card de autenticação; 3. O usuário digita e-mail/usuário e senha; 4. O usuário clica em "Entrar"; 5. O sistema valida as informações, emite mensagem de boas-vindas e redireciona automaticamente para a vitrine (/inicio).',
        'cenario_alternativo': 'Se algum campo obrigatório for deixado em branco, o sistema impede a submissão e emite um alerta orientando o preenchimento.'
    },
    {
        'id': 'UC04',
        'nome': 'Comprar Jogo / Adicionar ao Carrinho',
        'ator': 'Cliente Autenticado',
        'comunicacao': 'Clique no botão "Comprar" posicionado no rodapé do card de produto.',
        'cenario_principal': '1. O cliente localiza o jogo desejado na vitrine; 2. O cliente clica em "Comprar"; 3. O sistema registra a intenção de compra e conduz o cliente para o fluxo de checkout.',
        'cenario_alternativo': 'Caso o usuário não esteja logado, o sistema pode orientá-lo a efetuar o login prévio antes de concluir o pedido.'
    },
    {
        'id': 'UC05',
        'nome': 'Navegar pelo Menu e Rodapé',
        'ator': 'Usuário / Gamer (Visitante) e Cliente Autenticado',
        'comunicacao': 'Clique nos links de navegação da barra superior e do rodapé.',
        'cenario_principal': '1. O usuário clica em qualquer item de menu (ex.: "Produtos" ou "Login"); 2. O Angular Router interpreta a URL e renderiza o componente correspondente instantaneamente sem recarregar o navegador.',
        'cenario_alternativo': 'Se a rota não for encontrada, o sistema redireciona por padrão para /inicio.'
    }
]

for uc in ucs:
    p_uc = doc.add_paragraph()
    p_uc.paragraph_format.space_before = Pt(8)
    p_uc.paragraph_format.space_after = Pt(2)
    r_uc = p_uc.add_run(f'Caso de Uso {uc["id"]}: {uc["nome"]}')
    r_uc.font.bold = True
    r_uc.font.color.rgb = COLOR_PRIMARY
    r_uc.font.size = Pt(11)
    
    add_body_p(doc, uc['ator'], '• Ator(es): ', 2)
    add_body_p(doc, uc['comunicacao'], '• Comunicação: ', 2)
    add_body_p(doc, uc['cenario_principal'], '• Cenário Principal: ', 2)
    add_body_p(doc, uc['cenario_alternativo'], '• Cenário Alternativo: ', 6)

# 11. DEFINIÇÕES, ACRÔNIMOS E ABREVIAÇÕES
add_section_heading(doc, 11, 'Definições, Acrônimos e Abreviações')
def_table = doc.add_table(rows=9, cols=2)
def_table.alignment = WD_TABLE_ALIGNMENT.CENTER
def_headers = ['Termo / Sigla', 'Definição']
for col_idx, text in enumerate(def_headers):
    cell = def_table.rows[0].cells[col_idx]
    cell.text = text
    set_cell_background(cell, '3F51B5')
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    cell.paragraphs[0].runs[0].font.size = Pt(9.5)

defs_data = [
    ('SPA', 'Single Page Application (Aplicação de Página Única). Aplicação web que carrega uma única página HTML e atualiza o conteúdo dinamicamente conforme o usuário interage.'),
    ('Angular', 'Plataforma e framework open-source mantido pelo Google para desenvolvimento de aplicações web client-side em TypeScript.'),
    ('Angular Material', 'Biblioteca de componentes de interface de usuário (UI) baseada nas especificações do Google Material Design.'),
    ('Bootstrap', 'Framework front-end para criação de layouts responsivos, grids e componentes de interface web.'),
    ('API', 'Application Programming Interface (Interface de Programação de Aplicações). Conjunto de regras e protocolos para comunicação entre sistemas.'),
    ('REST', 'Representational State Transfer. Padrão arquitetural para desenvolvimento de serviços web baseado no protocolo HTTP.'),
    ('LGPD', 'Lei Geral de Proteção de Dados Pessoais (Lei nº 13.709/2018), que regulamenta o tratamento de dados pessoais no Brasil.'),
    ('UML', 'Unified Modeling Language (Linguagem de Modelagem Unificada). Padrão visual para especificação e documentação de sistemas de software.')
]
for row_idx, data in enumerate(defs_data, start=1):
    for col_idx, val in enumerate(data):
        cell = def_table.rows[row_idx].cells[col_idx]
        cell.text = val
        cell.paragraphs[0].runs[0].font.size = Pt(9.0)
        set_cell_background(cell, 'F5F7FB' if row_idx % 2 == 1 else 'FFFFFF')
        set_cell_margins(cell, 45, 45, 70, 70)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# 12. REFERÊNCIAS
add_section_heading(doc, 12, 'Referências')
add_bullet(doc, 'ANGULAR. Angular Documentation and Developer Guides. Disponível em: <https://angular.io/docs>. Acesso em: 2026.', '1. ')
add_bullet(doc, 'ANGULAR MATERIAL. Component UI Framework for Angular. Disponível em: <https://material.angular.io>. Acesso em: 2026.', '2. ')
add_bullet(doc, 'BOOTSTRAP. The most popular HTML, CSS, and JS library. Disponível em: <https://getbootstrap.com>. Acesso em: 2026.', '3. ')
add_bullet(doc, 'PRESSMAN, Roger S.; MAXIM, Bruce R. Engenharia de Software: Uma Abordagem Profissional. 8. ed. Porto Alegre: AMGH, 2016.', '4. ')
add_bullet(doc, 'SOMMERVILLE, Ian. Engenharia de Software. 10. ed. São Paulo: Pearson Education do Brasil, 2019.', '5. ')
add_bullet(doc, 'MATERIAL DIDÁTICO DO CURSO. Framework de Consumo de API - Situação de Aprendizagem 2 (SP2).', '6. ')

output_path = 'Framework_Consumo_API_SP2 MI_Framework_API_Documentacao de sistema.docx'
doc.save(output_path)
print(f'Document successfully generated at {output_path}')
